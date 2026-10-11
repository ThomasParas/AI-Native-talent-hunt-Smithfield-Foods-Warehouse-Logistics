"""Score suppliers against tool/criteria.yaml and emit JSON + XLSX scorecards.

Usage:  uv run --with pyyaml --with openpyxl python tool/score.py [out_dir]
Each tool/suppliers/*.yaml file is one supplier. Add a file to score a new supplier.
"""
import json
import sys
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parent
CRIT = yaml.safe_load((ROOT / "criteria.yaml").read_text())
MULT = CRIT["evidence_multiplier"]
CRITERIA = [c | {"group": g["name"]} for g in CRIT["groups"] for c in g["criteria"]]
assert sum(c["w"] for c in CRITERIA) == 100, "weights must sum to 100"


def gate_status(sup, scen):
    """FAIL if any gate fails, CONDITIONAL if any unverified, else PASS."""
    vals = [g.get(scen, "unverified") for g in sup["gates"].values()]
    if "fail" in vals:
        return "FAIL"
    if "unverified" in vals or "claimed" in vals:
        return "CONDITIONAL"
    return "PASS"


def score(sup):
    rows, raw, adj, known_w = [], 0.0, 0.0, 0
    for c in CRITERIA:
        s = sup["scores"].get(c["id"], {"score": 0, "evidence": "unverified", "note": "No information"})
        ev = s.get("evidence", "unverified")
        r = c["w"] * s["score"] / 5
        a = r * MULT[ev]
        raw += r if ev != "unverified" else 0
        adj += a
        known_w += c["w"] if ev != "unverified" else 0
        rows.append({**c, **s, "evidence": ev, "raw": round(r, 2), "adj": round(a, 2)})
    gaps = [r for r in rows if r["evidence"] == "unverified" or r["score"] <= 2]
    return {
        "name": sup["name"],
        "short": sup["short"],
        "rows": rows,
        "raw": round(raw, 1),
        "adjusted": round(adj, 1),
        "evidence_coverage": known_w,
        "gates": {k: gate_status(sup, k) for k in CRIT["scenarios"]},
        "gate_detail": sup["gates"],
        "gaps": gaps,
    }


def load_suppliers():
    files = sorted((ROOT / "suppliers").glob("*.yaml"))
    return [yaml.safe_load(f.read_text()) for f in files]


def write_xlsx(results, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter as L

    wb = Workbook()
    ws = wb.active
    ws.title = "Scorecard"
    head = ["Group", "ID", "Criterion", "Weight", "Scale (0-5)"]
    for r in results:
        head += [f"{r['short']} score", f"{r['short']} evidence", f"{r['short']} weighted"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
    mult = {"verified": 1, "claimed": 0.7, "unverified": 0}
    for i, c in enumerate(CRITERIA):
        row = [c["group"], c["id"], c["name"], c["w"], c["scale"]]
        n = i + 2
        for j, r in enumerate(results):
            col = 6 + j * 3
            s = r["rows"][i]
            sc, ev = L(col), L(col + 1)
            row += [s["score"], s["evidence"],
                    f'=D{n}*{sc}{n}/5*IF({ev}{n}="verified",1,IF({ev}{n}="claimed",0.7,0))']
        ws.append(row)
    last = len(CRITERIA) + 1
    tot = ["", "", "TOTAL (risk-adjusted, /100)", f"=SUM(D2:D{last})", ""]
    for j, _ in enumerate(results):
        col = L(8 + j * 3)
        tot += ["", "", f"=SUM({col}2:{col}{last})"]
    ws.append(tot)
    for c in ws[last + 1]:
        c.font = Font(bold=True)
    ws.column_dimensions["C"].width = 42
    ws.column_dimensions["E"].width = 60
    g = wb.create_sheet("Gates")
    g.append(["Gate", "Test"] + [f"{r['short']} {s}" for r in results for s in CRIT["scenarios"]])
    for gate in CRIT["gates"]:
        line = [f"{gate['id']} {gate['name']}", gate["test"]]
        for r in results:
            d = r["gate_detail"].get(gate["id"], {})
            line += [d.get(s, "unverified") for s in CRIT["scenarios"]]
        g.append(line)
    g.column_dimensions["B"].width = 70
    wb.save(path)


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else ROOT.parent / "build")
    out.mkdir(parents=True, exist_ok=True)
    results = [score(s) for s in load_suppliers()]
    (out / "scores.json").write_text(json.dumps({"criteria": CRIT, "results": results}, indent=1, ensure_ascii=False))
    write_xlsx(results, out / "supplier_scorecard.xlsx")
    for r in results:
        print(f"{r['short']:<12} adj={r['adjusted']:>5}/100 raw={r['raw']:>5} evidence={r['evidence_coverage']}% gates={r['gates']}")
